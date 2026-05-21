#!/usr/bin/env python3
"""Generate a .pptx file from a JSON slide definition.

Usage:
    python3 generate_ppt.py slides.json [output.pptx]

JSON format:
{
  "title": "Presentation Title",
  "slides": [
    {
      "layout": "title",           // "title" | "content" | "two_column"
      "title": "Slide Title",
      "subtitle": "Optional Subtitle"
    },
    {
      "layout": "content",
      "title": "Slide Title",
      "bullets": ["Point 1", "Point 2", "Point 3"]
    },
    {
      "layout": "two_column",
      "title": "Slide Title",
      "left": ["Left Point 1", "Left Point 2"],
      "right": ["Right Point 1", "Right Point 2"]
    }
  ]
}
"""

import json
import sys
import os

def install_pptx():
    """Install python-pptx if not available."""
    import subprocess
    subprocess.check_call(
        [sys.executable, "-m", "pip", "install", "python-pptx", "-q"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
except ImportError:
    print("Installing python-pptx...")
    install_pptx()
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


# Theme colors
COLOR_PRIMARY = RGBColor(0x1A, 0x73, 0xE8)
COLOR_DARK = RGBColor(0x20, 0x2A, 0x44)
COLOR_LIGHT_BG = RGBColor(0xF5, 0xF7, 0xFA)
COLOR_GRAY = RGBColor(0x6B, 0x7B, 0x8D)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_ACCENT = RGBColor(0x34, 0xA8, 0x53)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def add_title_slide(prs, slide_data):
    """Create a title slide with large centered text."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_PRIMARY

    # Title
    title_text = slide_data.get("title", "")
    if title_text:
        txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(2))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER

    # Subtitle
    subtitle_text = slide_data.get("subtitle", "")
    if subtitle_text:
        txBox = slide.shapes.add_textbox(Inches(2), Inches(4.5), Inches(9.3), Inches(1.2))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = subtitle_text
        p.font.size = Pt(22)
        p.font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        p.alignment = PP_ALIGN.CENTER


def add_content_slide(prs, slide_data):
    """Create a content slide with title and bullet points."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Top accent bar
    bar = slide.shapes.add_shape(
        1,  # Rectangle
        Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.08)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    # Title
    title_text = slide_data.get("title", "")
    if title_text:
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK

    # Bullets
    bullets = slide_data.get("bullets", [])
    if bullets:
        txBox = slide.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(5.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(bullets):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(20)
            p.font.color.rgb = COLOR_DARK
            p.space_after = Pt(12)
            p.level = 0
            # Bullet character
            p.text = "  " + bullet
            pPr = p._pPr
            if pPr is None:
                from pptx.oxml.ns import qn
                pPr = p._p.get_or_add_pPr()
            from pptx.oxml.ns import qn
            buChar = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            # Remove existing bullets
            for existing in pPr.findall(qn("a:buChar")):
                pPr.remove(existing)
            pPr.append(buChar)


def add_two_column_slide(prs, slide_data):
    """Create a two-column content slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Top accent bar
    bar = slide.shapes.add_shape(
        1,  # Rectangle
        Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.08)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    # Title
    title_text = slide_data.get("title", "")
    if title_text:
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK

    # Left column
    left_items = slide_data.get("left", [])
    if left_items:
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(5.6), Inches(5.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        for i, item in enumerate(left_items):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = "  " + item
            p.font.size = Pt(18)
            p.font.color.rgb = COLOR_DARK
            p.space_after = Pt(8)

    # Right column
    right_items = slide_data.get("right", [])
    if right_items:
        txBox = slide.shapes.add_textbox(Inches(6.8), Inches(2.0), Inches(5.6), Inches(5.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        for i, item in enumerate(right_items):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = "  " + item
            p.font.size = Pt(18)
            p.font.color.rgb = COLOR_DARK
            p.space_after = Pt(8)


def generate_ppt(slides_json, output_path):
    """Generate a .pptx file from slide JSON data."""
    if isinstance(slides_json, str):
        with open(slides_json, "r", encoding="utf-8") as f:
            data = json.load(f)
    else:
        data = slides_json

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slides = data.get("slides", [])
    for slide_data in slides:
        layout = slide_data.get("layout", "content")
        if layout == "title":
            add_title_slide(prs, slide_data)
        elif layout == "two_column":
            add_two_column_slide(prs, slide_data)
        else:
            add_content_slide(prs, slide_data)

    prs.save(output_path)
    return output_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 generate_ppt.py <slides.json> [output.pptx]")
        sys.exit(1)

    json_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else None

    if not out_path:
        base = os.path.splitext(os.path.basename(json_path))[0]
        out_path = os.path.join(os.path.dirname(json_path), base + ".pptx")

    result = generate_ppt(json_path, out_path)
    print(f"PPT generated: {result}")
